"""Salidas del reporte confirmado: PDF para el director y filas para la planilla de seguimiento."""
from pathlib import Path
from typing import TYPE_CHECKING

from fpdf import FPDF
from openpyxl import Workbook, load_workbook
from PIL import Image

from . import config, cronograma, db

if TYPE_CHECKING:
    from .reporte import Borrador

NIVELES = {"critico": "CRÍTICO", "rojo": "ROJO", "amarillo": "AMARILLO", "verde": "VERDE"}
TERRENO = {"seco": "seco", "humedo": "húmedo", "anegado": "anegado"}
TIPOS_NOTA = {
    "orden_de_servicio": "Orden de servicio", "nota_de_pedido": "Nota de pedido",
    "observacion_dt": "Observación del DT", "comunicacion_comitente": "Comunicación con el comitente",
    "acta": "Acta",
}
REEMPLAZOS = {"–": "-", "—": "-", "•": "-", "“": '"', "”": '"', "‘": "'", "’": "'", "…": "...", " ": " "}


def _t(texto) -> str:
    """Las fuentes base del PDF son latin-1: reemplaza los caracteres que no entran."""
    s = "" if texto is None else str(texto)
    for a, b in REEMPLAZOS.items():
        s = s.replace(a, b)
    return s.encode("latin-1", "replace").decode("latin-1")


def _fecha(iso: str) -> str:
    return f"{iso[8:10]}/{iso[5:7]}/{iso[0:4]}" if iso and len(iso) >= 10 else (iso or "")


class PDF(FPDF):
    titulo = ""

    def header(self):
        self.set_font("Helvetica", "B", 9)
        self.set_text_color(110)
        self.cell(0, 6, _t(self.titulo), new_x="LMARGIN", new_y="NEXT")
        self.set_text_color(0)
        self.ln(2)

    def footer(self):
        self.set_y(-12)
        self.set_font("Helvetica", "", 8)
        self.set_text_color(130)
        self.cell(0, 6, f"Página {self.page_no()}/{{nb}}", align="R")


def _seccion(pdf: PDF, titulo: str):
    pdf.ln(3)
    pdf.set_font("Helvetica", "B", 11)
    pdf.set_fill_color(235, 235, 235)
    pdf.cell(0, 7, _t(titulo), fill=True, new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("Helvetica", "", 9)
    pdf.ln(1)


def _parrafo(pdf: PDF, texto: str):
    pdf.multi_cell(0, 5, _t(texto), new_x="LMARGIN", new_y="NEXT")


def _tabla(pdf: PDF, encabezado: list[str], filas: list[list], anchos: tuple):
    if not filas:
        _parrafo(pdf, "Sin registros.")
        return
    pdf.set_font("Helvetica", "", 8.5)
    with pdf.table(col_widths=anchos, text_align="LEFT", line_height=4.6) as tabla:
        fila = tabla.row()
        for c in encabezado:
            fila.cell(_t(c))
        for f in filas:
            fila = tabla.row()
            for c in f:
                fila.cell(_t(c))
    pdf.set_font("Helvetica", "", 9)


def _fotos(pdf: PDF, fotos):
    ancho_max, alto_max, gap = 88, 66, 6
    x0 = pdf.l_margin
    col = 0
    y_fila = pdf.get_y()
    alto_fila = 0
    for f in fotos:
        path = Path(f["path"])
        if not path.exists():
            continue
        with Image.open(path) as img:
            w_px, h_px = img.size
        w = ancho_max
        h = w * h_px / w_px
        if h > alto_max:
            h = alto_max
            w = h * w_px / h_px
        if col == 0:
            if y_fila + alto_max + 14 > pdf.h - pdf.b_margin:
                pdf.add_page()
                y_fila = pdf.get_y()
            alto_fila = 0
        x = x0 + col * (ancho_max + gap)
        pdf.image(str(path), x=x, y=y_fila, w=w, h=h)
        pdf.set_xy(x, y_fila + h + 1)
        pdf.set_font("Helvetica", "", 7.5)
        epigrafe = f"#{f['id']} {f['hora'] or ''} - {f['id_tarea'] or 'sin tarea'}: {f['descripcion'] or ''}"
        pdf.multi_cell(ancho_max, 3.6, _t(epigrafe[:220]))
        alto_fila = max(alto_fila, pdf.get_y() - y_fila)
        col += 1
        if col == 2:
            col = 0
            y_fila += alto_fila + 4
    pdf.set_font("Helvetica", "", 9)
    pdf.set_xy(x0, y_fila + (alto_fila + 4 if col else 0))


def generar_pdf(b: "Borrador") -> Path:
    d = b.datos
    obra = b.obra
    real, previsto = cronograma.avance_obra(obra.id, b.tareas, b.fecha)
    fotos = db.fotos_de_reporte(b.id)
    alertas = d["proximos"]["pedidos"]

    pdf = PDF(format="A4")
    pdf.titulo = f"Reporte diario de obra N° {b.numero} - {obra.nombre}"
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_page()

    # 1. Encabezado
    pdf.set_font("Helvetica", "B", 16)
    pdf.cell(0, 9, _t(f"Reporte diario N° {b.numero}"), new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("Helvetica", "", 10)
    _parrafo(pdf, f"Obra: {obra.nombre} - {obra.direccion}")
    _parrafo(pdf, f"Fecha: {_fecha(b.fecha.isoformat())}    DT: {b.dt_nombre}    "
                  f"Avance de obra: {real}% real / {previsto}% previsto")

    # 2. Resumen
    _seccion(pdf, "Resumen")
    _parrafo(pdf, d.get("resumen") or "-")

    # 3. Alertas
    _seccion(pdf, "Alertas")
    lineas = []
    for a in alertas:
        if a["nivel"] in ("critico", "rojo"):
            lineas.append(f"[{NIVELES[a['nivel']]}] Pedir {a['material']} ({a['cantidad']:g} {a['unidad']}) "
                          f"para {a['tarea']}: fecha límite {_fecha(a['fecha_limite_pedido'])}")
    for av in d["avances"]:
        if av.get("desvio") is not None and av["desvio"] <= -10:
            lineas.append(f"[DESVÍO] {av['tarea']}: {av['pct_real']:g}% real vs {av['pct_previsto']:g}% previsto")
    for n in db.notas_abiertas(obra.id):
        if n["plazo"] and n["plazo"] < b.fecha.isoformat():
            lineas.append(f"[NOTA VENCIDA] {TIPOS_NOTA.get(n['tipo'], n['tipo'])} N° {n['numero']}: {n['asunto']}")
    _parrafo(pdf, "\n".join(f"- {x}" for x in lineas) if lineas else "Sin alertas.")

    # 4. Condiciones y jornada
    _seccion(pdf, "Condiciones y jornada")
    c, j = d["condiciones"], d["jornada"]
    _parrafo(pdf, f"Clima: mañana {c.get('clima_manana', '-')}, tarde {c.get('clima_tarde', '-')}. "
                  f"Temp. {c.get('temp_min', '-')} a {c.get('temp_max', '-')} °C. Lluvia {c.get('lluvia_mm', 0)} mm. "
                  f"Terreno: {TERRENO.get(c.get('terreno'), '-')}.")
    jornada = f"Jornada: {j.get('tipo', '-').replace('_', ' ')}"
    if j.get("horas_perdidas"):
        jornada += f". Horas perdidas: {j['horas_perdidas']:g} ({j.get('causa', 'sin causa')})"
    _parrafo(pdf, jornada + ".")

    # 5. Avance
    _seccion(pdf, "Avance real contra previsto")
    _tabla(pdf, ["Tarea", "Descripción", "% real", "% prev.", "Desvío"], [
        [f"{a['id_tarea']} {a['tarea']}", a["descripcion"], f"{a['pct_real']:g}",
         "-" if a.get("pct_previsto") is None else f"{a['pct_previsto']:g}",
         "-" if a.get("desvio") is None else f"{a['desvio']:+g}"]
        for a in d["avances"]
    ], (50, 80, 18, 18, 18))

    # 6. Personal
    _seccion(pdf, "Personal")
    _tabla(pdf, ["Empresa", "Categoría", "Cantidad", "Horas"], [
        [p["empresa"], p["categoria"], p["cantidad"], p.get("horas", "")] for p in d["personal"]
    ], (70, 60, 25, 25))

    # 7. Materiales
    _seccion(pdf, "Materiales")
    m = d["materiales"]
    filas = [["Ingreso", x["material"], f"{x['cantidad']:g} {x['unidad']}", x.get("proveedor", ""), x.get("remito", "")]
             for x in m["ingresos"]]
    filas += [["Faltante", x["material"], f"{x['cantidad']:g} {x['unidad']}", "", ""] for x in m["faltantes"]]
    filas += [["Stock", x["material"], f"{x['cantidad']:g} {x['unidad']}", "", ""] for x in m["stock"]]
    _tabla(pdf, ["Tipo", "Material", "Cantidad", "Proveedor", "Remito"], filas, (22, 60, 30, 45, 25))
    if d["equipos"]:
        _parrafo(pdf, "Equipos: " + "; ".join(f"{e['equipo']} ({e['estado']})" for e in d["equipos"]))

    # 8. Notas de obra
    _seccion(pdf, "Notas de obra abiertas")
    notas = db.notas_abiertas(obra.id)
    if not notas:
        _parrafo(pdf, "Sin notas abiertas.")
    for n in notas:
        pdf.set_font("Helvetica", "B", 9)
        _parrafo(pdf, f"{TIPOS_NOTA.get(n['tipo'], n['tipo'])} N° {n['numero']} - {_fecha(n['fecha'])} - "
                      f"{n['emisor']} a {n['destinatario']}: {n['asunto']}"
                      + (f" (plazo {_fecha(n['plazo'])})" if n["plazo"] else ""))
        pdf.set_font("Helvetica", "I", 9)
        _parrafo(pdf, f'"{n["texto"]}"')
        pdf.set_font("Helvetica", "", 9)
        pdf.ln(1)

    # 9. Higiene y seguridad
    _seccion(pdf, "Higiene y seguridad")
    h = d["hys"]
    _parrafo(pdf, f"Uso de EPP: {'sí' if h.get('epp_ok') else 'no'}. Visitas: {h.get('visitas') or 'no hubo'}. "
                  f"Incidentes: {h.get('incidentes') or 'no hubo'}.")

    # 10. Próximos pasos
    _seccion(pdf, f"Próximos pasos ({_fecha(d['proximos']['fecha'])})")
    tareas = d["proximos"]["tareas"]
    _parrafo(pdf, "Tareas: " + ("; ".join(f"{t['id_tarea']} {t['tarea']}" for t in tareas) if tareas else "sin tareas previstas"))
    if alertas:
        _parrafo(pdf, "Pedidos a gestionar:\n" + "\n".join(
            f"- [{NIVELES[a['nivel']]}] {a['material']} {a['cantidad']:g} {a['unidad']} "
            f"({a['proveedor'] or 'sin proveedor'}), pedir antes del {_fecha(a['fecha_limite_pedido'])}"
            for a in alertas))

    # 11. Fotos
    _seccion(pdf, f"Fotos ({len(fotos)})")
    _fotos(pdf, fotos)

    pdf.ln(4)
    pdf.set_font("Helvetica", "I", 8.5)
    _parrafo(pdf, f"Reporte confirmado por {b.dt_nombre} el {d['confirmado_en']}.")

    destino = config.SALIDAS_DIR / obra.id
    destino.mkdir(parents=True, exist_ok=True)
    path = destino / f"reporte_{b.fecha.isoformat()}_N{b.numero:03d}.pdf"
    pdf.output(str(path))
    return path


# --- planilla de seguimiento -------------------------------------------------

HOJAS = {
    "Diario": ["fecha", "obra", "n_reporte", "dt", "clima_manana", "clima_tarde", "lluvia_mm", "jornada",
               "horas_perdidas", "causa", "personal_total", "avance_real_pct", "avance_previsto_pct",
               "desvio_pct", "alertas_rojas", "notas_abiertas", "link_pdf"],
    "Avance": ["fecha", "obra", "id_tarea", "rubro", "tarea", "pct_real", "pct_previsto", "desvio_pct", "fotos"],
    "Materiales": ["fecha", "obra", "material", "unidad", "stock", "ingreso", "remito", "proveedor",
                   "fecha_limite_pedido", "nivel_alerta", "estado_pedido"],
    "Notas": ["n_nota", "tipo", "obra", "fecha", "emisor", "destinatario", "asunto", "plazo", "estado",
              "impacto_materiales"],
}


def _libro():
    if config.PLANILLA_PATH.exists():
        return load_workbook(config.PLANILLA_PATH)
    wb = Workbook()
    wb.remove(wb.active)
    for nombre, columnas in HOJAS.items():
        ws = wb.create_sheet(nombre)
        ws.append(columnas)
    return wb


def agregar_a_planilla(b: "Borrador"):
    d = b.datos
    obra = b.obra.nombre
    fecha = b.fecha.isoformat()
    real, previsto = cronograma.avance_obra(b.obra.id, b.tareas, b.fecha)
    alertas = d["proximos"]["pedidos"]
    fotos = db.fotos_de_reporte(b.id)
    wb = _libro()

    j = d["jornada"]
    wb["Diario"].append([
        fecha, obra, b.numero, b.dt_nombre, d["condiciones"].get("clima_manana"),
        d["condiciones"].get("clima_tarde"), d["condiciones"].get("lluvia_mm"), j.get("tipo"),
        j.get("horas_perdidas", 0), j.get("causa", ""), sum(p["cantidad"] for p in d["personal"]),
        real, previsto, round(real - previsto, 1),
        sum(1 for a in alertas if a["nivel"] in ("rojo", "critico")),
        len(db.notas_abiertas(b.obra.id)), d.get("pdf", ""),
    ])
    for a in d["avances"]:
        n_fotos = sum(1 for f in fotos if f["id_tarea"] == a["id_tarea"])
        wb["Avance"].append([fecha, obra, a["id_tarea"], a["rubro"], a["tarea"], a["pct_real"],
                             a.get("pct_previsto"), a.get("desvio"), n_fotos])
    m = d["materiales"]
    for x in m["ingresos"]:
        wb["Materiales"].append([fecha, obra, x["material"], x["unidad"], None, x["cantidad"],
                                 x.get("remito", ""), x.get("proveedor", ""), None, None, "recibido"])
    for x in m["stock"]:
        wb["Materiales"].append([fecha, obra, x["material"], x["unidad"], x["cantidad"], None,
                                 "", "", None, None, None])
    for a in alertas:
        wb["Materiales"].append([fecha, obra, a["material"], a["unidad"], a["stock_declarado"], None, "",
                                 a["proveedor"], a["fecha_limite_pedido"], a["nivel"], "pendiente"])
    for n in db.notas_de_reporte(b.id):
        wb["Notas"].append([n["numero"], n["tipo"], obra, n["fecha"], n["emisor"], n["destinatario"],
                            n["asunto"], n["plazo"], n["estado"], bool(n["impacto_materiales"])])
    wb.save(config.PLANILLA_PATH)
